"""Extrae los datos adjuntos sin modificar los libros ni la plantilla oficial."""

from pathlib import Path
import argparse
import collections
import json
import re
import unicodedata

import openpyxl


RAIZ = Path(__file__).resolve().parents[1]
PROVINCIAS = {
    "1": "San José", "2": "Alajuela", "3": "Cartago", "4": "Heredia",
    "5": "Guanacaste", "6": "Puntarenas", "7": "Limón",
}


def texto(valor):
    return re.sub(r"\s+", " ", str(valor or "")).strip()


def clave(valor):
    # Permite cotejar códigos sin alterar su presentación original.
    return str(int(valor)) if isinstance(valor, (int, float)) else texto(valor)


def entero(valor):
    # El Excel mezcla números y números guardados como texto (ej. Sarapiquí: "10").
    v = str(valor).strip()
    return int(float(v)) if re.fullmatch(r"\d+(?:\.0+)?", v) else None


def sin_acentos(valor):
    return "".join(c for c in unicodedata.normalize("NFD", valor.casefold())
                   if not unicodedata.combining(c))


def partir_tramo(descripcion):
    """Separa DE/A solo cuando hay un delimitador inequívoco fuera de paréntesis."""
    nivel, separadores = 0, []
    for i, c in enumerate(descripcion):
        if c == "(":
            nivel += 1
        elif c == ")":
            nivel = max(0, nivel - 1)
        elif c in "-–—" and nivel == 0:
            separadores.append(i)
    if len(separadores) == 1:
        i = separadores[0]
        inicio, fin = descripcion[:i].strip(), descripcion[i + 1:].strip()
        if inicio and fin:
            return inicio, fin, False
    # Nunca se reparte una descripción ambigua como si estuviera confirmada.
    return "", "", True


def codigos_rtf(ruta):
    """Relaciona códigos y secciones que aparecen explícitamente en el RTF."""
    raw = ruta.read_text(encoding="cp1252")
    parrafos = []
    for m in re.finditer(r"\{\\pard\b.*?\\par\}", raw, re.S):
        partes = re.findall(r"\\overlay\s+((?:(?!\\overlay|\\par).)*)", m.group(), re.S)
        if partes:
            t = partes[-1]
            t = re.sub(r"\\'([a-fA-F0-9]{2})", lambda h: bytes.fromhex(h[1]).decode("cp1252"), t)
            t = re.sub(r"\\[a-z]+-?\d* ?", "", t)
            parrafos.append(texto(t))
    relacion = {}
    for i, t in enumerate(parrafos):
        if re.fullmatch(r"4-\d{4}-\d{2}", t):
            # La sección de control sigue al código en cada fila del informe.
            for candidato in parrafos[i + 1:i + 5]:
                if re.fullmatch(r"\d{5}", candidato):
                    relacion[candidato] = t
                    break
    return relacion


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--referencias", type=Path, default=RAIZ / "referencias")
    args = parser.parse_args()
    libro = openpyxl.load_workbook(args.referencias / "red-vial-nacional-2024.xlsx", read_only=True, data_only=True)
    relacion = codigos_rtf(args.referencias / "red-nacional-heredia.rtf")

    cantones = {}
    for fila in libro["Longitud Por Provincia"].iter_rows(values_only=True):
        if isinstance(fila[1], (int, float)) and 100 <= fila[1] < 800 and isinstance(fila[2], str):
            codigo = str(int(fila[1]))
            cantones[codigo] = texto(fila[2])

    caminos, omitidas = [], []

    def agregar(fila, hoja, n, seccion, provincia, canton, ruta, longitud, descripcion):
        seccion, provincia, canton = entero(seccion), entero(provincia), entero(canton)
        if seccion is None or provincia is None:
            return
        if str(provincia) not in PROVINCIAS or canton is None:
            return
        if not isinstance(longitud, (int, float)) or longitud <= 0 or not descripcion:
            omitidas.append({"hoja": hoja, "fila": n, "motivo": "sin longitud positiva o descripción"})
            return
        sec, prov, cant = clave(seccion), clave(provincia), f"{int(canton):02d}"
        canton_nombre = cantones.get(prov + cant)
        if not canton_nombre:
            omitidas.append({"hoja": hoja, "fila": n, "motivo": f"cantón {prov + cant} sin nombre en el resumen"})
            return
        desc = texto(descripcion)
        inicio, fin, revisar = partir_tramo(desc)
        # El RTF confirma el formato de código en las secciones relacionadas.
        # Las demás conservan la sección exacta del Excel; no se inventa un código.
        codigo = relacion.get(sec, sec)
        caminos.append({
            "id": f"N:{hoja}:{sec}:{n}", "type": "N", "code": codigo, "section": sec,
            "route": clave(ruta), "lengthKm": round(float(longitud), 6),
            "from": inicio, "to": fin, "description": desc,
            "provinceCode": prov, "provinceName": PROVINCIAS[prov],
            "cantonCode": cant, "cantonName": canton_nombre,
            "network": "Travesía" if hoja == "Travesías" else "Ruta nacional",
            "needsReview": revisar,
            "source": {"file": "Red Vial Nacional 2024(1).xlsx", "sheet": hoja, "row": n},
            "codeSource": "red nacional heredia(1).rtf" if sec in relacion else "Sección del Excel 2024",
        })

    for n, f in enumerate(libro["RVN_0524"].iter_rows(min_row=2, values_only=True), 2):
        agregar(f, "RVN_0524", n, f[1], f[3], f[4], f[6], f[10], f[49])
    for n, f in enumerate(libro["Travesías"].iter_rows(min_row=9, values_only=True), 9):
        agregar(f, "Travesías", n, f[0], f[2], f[3], f[4], f[6], f[32])
    caminos.sort(key=lambda c: (int(c["provinceCode"]), int(c["section"]), c["id"]))
    auditoria = {
        "roadCount": len(caminos),
        "herediaCount": sum(c["provinceCode"] == "4" for c in caminos),
        "rtfCodeCount": sum(c["code"] != c["section"] for c in caminos),
        "needsReviewCount": sum(c["needsReview"] for c in caminos),
        "cantonCount": len(cantones), "skippedRows": omitidas,
        "duplicateSections": [s for s, total in collections.Counter(c["section"] for c in caminos).items() if total > 1],
    }
    salida = {"version": 1, "sourceYear": 2024, "provinces": PROVINCIAS, "cantons": cantones,
              "roads": caminos, "audit": auditoria}
    destino = RAIZ / "dist" / "data" / "caminos.json"
    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_text(json.dumps(salida, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(json.dumps(auditoria, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
